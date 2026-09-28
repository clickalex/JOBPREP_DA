"""
Build practice/excel/ShopKart_Excel_Practice.xlsx

    python practice/excel/build_workbook.py

Sheets
  Start Here   – instructions
  Orders       – Q4-2025 order lines (Excel Table `tblOrders`)   <- you work here
  Products     – product master (Excel Table `tblProducts`)
  Tasks        – 20 questions; type answers in the yellow cells and they self-grade
  Answer Key   – worked formulas for every task (don't peek too early!)
  Orders_Solved– the helper columns the answer key relies on

Expected answers are computed here with pandas and stored in a hidden column
of the Tasks sheet so each answer cell can grade itself.
"""
from __future__ import annotations

import sys

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Windows consoles/pipes default to cp1252, which can't print ✓ ✗ ₹ — force UTF-8.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
DATA = HERE.parents[1] / "data" / "clean"
OUT = HERE / "ShopKart_Excel_Practice.xlsx"

NAVY, YELLOW, GREY = "1F3864", "FFF2CC", "F2F2F2"
HFONT = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# --------------------------------------------------------------------------- data
def load():
    orders = pd.read_csv(DATA / "orders.csv", parse_dates=["order_ts"])
    items = pd.read_csv(DATA / "order_items.csv")
    customers = pd.read_csv(DATA / "customers.csv")
    products = pd.read_csv(DATA / "products.csv")
    o = orders[(orders.order_ts >= "2025-10-01") & (orders.order_ts < "2026-01-01")]
    df = (items.merge(o, on="order_id")
               .merge(customers[["customer_id", "city", "region", "acquisition_channel"]], on="customer_id"))
    df["order_date"] = df["order_ts"].dt.normalize()
    df = df.sort_values(["order_date", "order_id", "order_item_id"])
    cols = ["order_id", "order_date", "customer_id", "city", "region", "acquisition_channel", "device",
            "payment_method", "status", "product_id", "quantity", "unit_price", "discount"]
    df = df[cols].rename(columns={"acquisition_channel": "channel"}).reset_index(drop=True)
    prods = products[["product_id", "product_name", "category", "subcategory", "list_price", "unit_cost"]]
    return df, prods


# --------------------------------------------------------------------------- tasks
# (id, skill, question, answer type, hint, answer-key formula)
# Answer-key formulas reference the Orders_Solved sheet (helper columns N:R).
R = "Orders_Solved!"
N_ROWS = None  # filled at runtime


def rng(col):
    return f"{R}${col}$2:${col}${N_ROWS + 1}"


def tasks():
    # Orders columns: A order_id, B order_date, C customer_id, D city, E region, F channel, G device,
    #                 H payment_method, I status, J product_id, K quantity, L unit_price, M discount
    A, B, C, E, G, H, I, K, L, M = (rng(c) for c in "ABCEGHIKLM")
    CAT, REV, MON, SIZE, COST = (rng(c) for c in "NOPQR")
    return [
        (1, "XLOOKUP / VLOOKUP",
         "Add a Category column to the Orders table by looking up product_id in tblProducts. How many order lines are Fashion?",
         "num", "=XLOOKUP([@product_id], tblProducts[product_id], tblProducts[category]) then COUNTIF",
         f'=COUNTIF({CAT},"Fashion")'),
        (2, "Calculated column",
         "Add a Line Revenue column = quantity × unit_price − discount. What is the total Line Revenue across ALL lines (any status)?",
         "num", "=[@quantity]*[@unit_price]-[@discount], then SUM the column",
         f"=SUM({REV})"),
        (3, "SUMIFS (OR logic)",
         "Net revenue from VALID lines only (status Delivered or Shipped)?",
         "num", "SUMIFS can't do OR in one criterion – add two SUMIFS, or SUM(SUMIFS(..., {\"Delivered\",\"Shipped\"}))",
         f'=SUMIFS({REV},{I},"Delivered")+SUMIFS({REV},{I},"Shipped")'),
        (4, "SUMIFS (multiple criteria)",
         "Valid net revenue from Electronics in the South region?",
         "num", "Three criteria ranges: Category, region, status (x2 statuses)",
         f'=SUMIFS({REV},{CAT},"Electronics",{E},"South",{I},"Delivered")+SUMIFS({REV},{CAT},"Electronics",{E},"South",{I},"Shipped")'),
        (5, "COUNTIFS",
         "How many order lines were paid Cash on Delivery AND Cancelled?",
         "num", "COUNTIFS(range1, crit1, range2, crit2)",
         f'=COUNTIFS({H},"Cash on Delivery",{I},"Cancelled")'),
        (6, "Distinct count",
         "How many DISTINCT orders (order_id) are in the sheet? (Lines ≠ orders!)",
         "num", "Excel 365: =COUNTA(UNIQUE(tblOrders[order_id])). Older: =SUMPRODUCT(1/COUNTIF(rng,rng)) or Remove Duplicates on a copy",
         f"=SUMPRODUCT(1/COUNTIF({A},{A}))"),
        (7, "AVERAGEIFS",
         "Average unit_price of Beauty order lines? (2 dp)",
         "num", "AVERAGEIFS(avg_range, criteria_range, criteria)",
         f'=AVERAGEIFS({L},{CAT},"Beauty")'),
        (8, "MAXIFS",
         "Largest single Line Revenue in Home & Kitchen?",
         "num", "MAXIFS(max_range, criteria_range, criteria)  (Excel 2019+)",
         f'=_xlfn.MAXIFS({REV},{CAT},"Home & Kitchen")'),   # post-2007 functions need _xlfn. in the file
        (9, "Dates in criteria",
         "Valid net revenue for November 2025?",
         "num", "Use \">=\"&DATE(2025,11,1) and \"<\"&DATE(2025,12,1) as criteria on order_date",
         f'=SUMIFS({REV},{B},">="&DATE(2025,11,1),{B},"<"&DATE(2025,12,1),{I},"Delivered")'
         f'+SUMIFS({REV},{B},">="&DATE(2025,11,1),{B},"<"&DATE(2025,12,1),{I},"Shipped")'),
        (10, "IF",
         "Add an Order Size column: \"Large\" if Line Revenue ≥ 3000, else \"Small\". How many lines are Large?",
         "num", "=IF([@[Line Revenue]]>=3000,\"Large\",\"Small\")",
         f'=COUNTIF({SIZE},"Large")'),
        (11, "INDEX / MATCH",
         "What is the product_name of product_id 131? (use INDEX/MATCH on tblProducts)",
         "text", "=INDEX(tblProducts[product_name], MATCH(131, tblProducts[product_id], 0))",
         "=INDEX(Products!$B$2:$B$49,MATCH(131,Products!$A$2:$A$49,0))"),
        (12, "Ratios",
         "Return rate for Fashion = Returned lines ÷ (Delivered + Returned lines). Answer as a % (e.g. 12.3%).",
         "pct", "Two COUNTIFS divided; format as %",
         f'=COUNTIFS({CAT},"Fashion",{I},"Returned")/(COUNTIFS({CAT},"Fashion",{I},"Returned")+COUNTIFS({CAT},"Fashion",{I},"Delivered"))'),
        (13, "SUMPRODUCT",
         "Gross margin % on valid lines = (revenue − quantity × unit_cost) ÷ revenue. Answer as %.",
         "pct", "Look up unit_cost too, then SUMPRODUCT with a (status=\"Delivered\")+(status=\"Shipped\") mask",
         f'=(SUMPRODUCT({REV},({I}="Delivered")+({I}="Shipped"))-SUMPRODUCT({K},{COST},({I}="Delivered")+({I}="Shipped")))'
         f'/SUMPRODUCT({REV},({I}="Delivered")+({I}="Shipped"))'),
        (14, "PivotTable",
         "Which city generated the most valid net revenue? (type the city name)",
         "text", "Insert ▸ PivotTable: Rows = city, Values = Sum of Line Revenue, Filter = status. Sort descending.",
         "=INDEX('Answer Key'!$H$4:$H$18,MATCH(MAX('Answer Key'!$I$4:$I$18),'Answer Key'!$I$4:$I$18,0))"),
        (15, "PivotTable – % of total",
         "What share of valid net revenue came from the Mobile App? Answer as %.",
         "pct", "Pivot: Rows = device, Values = Sum of Line Revenue ▸ Show Values As ▸ % of Grand Total",
         f'=(SUMIFS({REV},{G},"Mobile App",{I},"Delivered")+SUMIFS({REV},{G},"Mobile App",{I},"Shipped"))/(SUMIFS({REV},{I},"Delivered")+SUMIFS({REV},{I},"Shipped"))'),
        (16, "Date functions",
         "Total discount given in December 2025 (any status)?",
         "num", "SUMIFS on discount with a date range – or add a Month column with =TEXT([@order_date],\"yyyy-mm\")",
         f'=SUMIFS({M},{MON},"2025-12")'),
        (17, "TEXT / WEEKDAY",
         "Which weekday had the most order lines? (e.g. Monday)",
         "text", "Add =TEXT([@order_date],\"dddd\") and pivot / COUNTIF it",
         "=INDEX('Answer Key'!$K$4:$K$10,MATCH(MAX('Answer Key'!$L$4:$L$10),'Answer Key'!$L$4:$L$10,0))"),
        (18, "Distinct count with a condition",
         "Average order value of valid orders = valid revenue ÷ number of distinct valid orders? (2 dp)",
         "num", "Distinct valid orders: =SUMPRODUCT(((status=\"Delivered\")+(status=\"Shipped\"))/COUNTIF(order_id,order_id))",
         f'=(SUMIFS({REV},{I},"Delivered")+SUMIFS({REV},{I},"Shipped"))/SUMPRODUCT((({I}="Delivered")+({I}="Shipped"))/COUNTIF({A},{A}))'),
        (19, "Distinct count + dates",
         "How many distinct customers placed an order in December 2025?",
         "num", "Excel 365: =COUNTA(UNIQUE(FILTER(customer_id, month=\"2025-12\"))). Older: SUMPRODUCT((month=\"2025-12\")/COUNTIFS(customer_id,customer_id,month,month))",
         f'=SUMPRODUCT(({MON}="2025-12")/COUNTIFS({C},{C},{MON},{MON}))'),
        (20, "Average of top N",
         "What is the average Line Revenue of the 10 largest lines (any status)? (2 dp)",
         "num", "=AVERAGE(LARGE(range,{1,2,3,4,5,6,7,8,9,10})) or SORT/TAKE in Excel 365",
         f"=SUMPRODUCT(LARGE({REV},ROW($1:$10)))/10"),
    ]


def expected(df, prods):
    d = df.merge(prods[["product_id", "product_name", "category", "unit_cost"]], on="product_id", how="left")
    d["rev"] = d.quantity * d.unit_price - d.discount
    valid = d.status.isin(["Delivered", "Shipped"])
    v = d[valid]
    month = d.order_date.dt.strftime("%Y-%m")
    fashion = d[(d.category == "Fashion") & d.status.isin(["Delivered", "Returned"])]
    return {
        1: int((d.category == "Fashion").sum()),
        2: d.rev.sum(),
        3: v.rev.sum(),
        4: v[(v.category == "Electronics") & (v.region == "South")].rev.sum(),
        5: int(((d.payment_method == "Cash on Delivery") & (d.status == "Cancelled")).sum()),
        6: int(d.order_id.nunique()),
        7: d[d.category == "Beauty"].unit_price.mean(),
        8: d[d.category == "Home & Kitchen"].rev.max(),
        9: v[month[valid] == "2025-11"].rev.sum(),
        10: int((d.rev >= 3000).sum()),
        11: prods.set_index("product_id").loc[131, "product_name"],
        12: (fashion.status == "Returned").mean(),
        13: (v.rev.sum() - (v.quantity * v.unit_cost).sum()) / v.rev.sum(),
        14: v.groupby("city").rev.sum().idxmax(),
        15: v[v.device == "Mobile App"].rev.sum() / v.rev.sum(),
        16: d[month == "2025-12"].discount.sum(),
        17: d.order_date.dt.day_name().value_counts().idxmax(),
        18: v.rev.sum() / v.order_id.nunique(),
        19: int(d[month == "2025-12"].customer_id.nunique()),
        20: d.rev.nlargest(10).mean(),
    }


# --------------------------------------------------------------------------- workbook helpers
def text_cell(ws, row, col, text):
    """Write literal text even if it starts with '=' (openpyxl would treat it as a formula)."""
    c = ws.cell(row=row, column=col)
    c.value = text
    c.data_type = "s"
    c.quotePrefix = True
    return c


def header(ws, row, values, widths=None):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.font, c.fill, c.alignment = HFONT, HFILL, Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def write_table(ws, df, name, widths, number_formats=None):
    header(ws, 1, list(df.columns), widths)
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
    for col, fmt in (number_formats or {}).items():
        j = list(df.columns).index(col) + 1
        for i in range(2, len(df) + 2):
            ws.cell(row=i, column=j).number_format = fmt
    ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "A2"


def build():
    global N_ROWS
    df, prods = load()
    N_ROWS = len(df)
    exp = expected(df, prods)
    wb = Workbook()

    # ---------------- Start Here
    ws = wb.active
    ws.title = "Start Here"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    lines = [
        ("ShopKart — Excel practice workbook", Font(bold=True, size=18, color=NAVY)),
        ("Data analyst interview prep · Q4 2025 order lines from the fictional ShopKart store", Font(italic=True, color="595959")),
        ("", None),
        ("HOW TO USE", Font(bold=True, size=12, color=NAVY)),
        ("1. Work on the Orders sheet. It is an Excel Table (tblOrders) — add your own columns to the right; formulas fill down automatically.", None),
        ("2. Products is the product master (tblProducts). You'll need it for Category, product_name and unit_cost.", None),
        ("3. Go to Tasks. Type each answer in the yellow cell — the Result column tells you if you're right (numbers are checked to 1 decimal place).", None),
        ("4. Percentages: type 12.3% (or 0.123). Text answers are not case-sensitive.", None),
        ("5. Stuck? The Answer Key sheet shows one working formula per task (it uses helper columns on Orders_Solved).", None),
        ("", None),
        ("DEFINITIONS", Font(bold=True, size=12, color=NAVY)),
        ("Line Revenue = quantity × unit_price − discount   (INR)", None),
        ("Valid line = status is Delivered or Shipped (Cancelled / Returned lines earn nothing)", None),
        ("One order can have several lines (rows) — be careful when counting orders vs lines.", None),
        ("", None),
        ("SKILLS COVERED", Font(bold=True, size=12, color=NAVY)),
        ("XLOOKUP · VLOOKUP · INDEX/MATCH · SUMIFS · COUNTIFS · AVERAGEIFS · MAXIFS · IF · TEXT · DATE · SUMPRODUCT · "
         "distinct counts · LARGE · PivotTables · % of total", None),
        ("", None),
        ("BONUS (not auto-graded)", Font(bold=True, size=12, color=NAVY)),
        ("• Build a PivotTable of valid revenue by Month × Category and add a PivotChart + a Slicer on region.", None),
        ("• Conditional formatting: highlight Cancelled rows in red using a formula rule: =$I2=\"Cancelled\".", None),
        ("• Build a one-page dashboard sheet: 4 KPI cards (Revenue, Orders, AOV, Return rate) + 2 charts.", None),
        ("• Use Power Query (Data ▸ Get Data ▸ From Text/CSV) to load data/raw/orders_raw.csv and clean the status / payment_method columns.", None),
    ]
    for i, (text, font) in enumerate(lines, start=2):
        c = ws.cell(row=i, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font

    # ---------------- Orders / Products
    ws_o = wb.create_sheet("Orders")
    write_table(ws_o, df, "tblOrders", [10, 12, 11, 12, 8, 15, 12, 17, 11, 10, 9, 10, 9],
                {"order_date": "yyyy-mm-dd", "unit_price": "#,##0.00", "discount": "#,##0.00"})
    ws_p = wb.create_sheet("Products")
    write_table(ws_p, prods, "tblProducts", [10, 28, 16, 16, 10, 10],
                {"list_price": "#,##0.00", "unit_cost": "#,##0.00"})

    # ---------------- Tasks
    ws_t = wb.create_sheet("Tasks")
    header(ws_t, 1, ["#", "Skill", "Question", "Your answer", "Result", "expected (hidden)", "Hint"],
           [4, 22, 70, 18, 14, 14, 70])
    ws_t.row_dimensions[1].height = 22
    task_list = tasks()
    for r, (tid, skill, q, kind, hint, _) in enumerate(task_list, start=2):
        ws_t.cell(row=r, column=1, value=tid)
        ws_t.cell(row=r, column=2, value=skill).font = Font(bold=True)
        ws_t.cell(row=r, column=3, value=q)
        ans = ws_t.cell(row=r, column=4)
        ans.fill = PatternFill("solid", fgColor=YELLOW)
        ans.border = BOX
        e = exp[tid]
        ws_t.cell(row=r, column=6, value=round(float(e), 6) if kind != "text" else e)
        if kind == "num":
            ans.number_format = "#,##0.00"
            check = f'=IF(D{r}="","",IF(ISNUMBER(D{r}),IF(ABS(D{r}-F{r})<0.05,"✅ Correct","❌ Try again"),"❌ Not a number"))'
        elif kind == "pct":
            ans.number_format = "0.0%"
            check = (f'=IF(D{r}="","",IF(ISNUMBER(D{r}),IF(OR(ABS(D{r}-F{r})<0.0005,ABS(D{r}/100-F{r})<0.0005),'
                     f'"✅ Correct","❌ Try again"),"❌ Not a number"))')
        else:
            check = f'=IF(D{r}="","",IF(LOWER(TRIM(D{r}))=LOWER(F{r}),"✅ Correct","❌ Try again"))'
        ws_t.cell(row=r, column=5, value=check)
        text_cell(ws_t, r, 7, hint).font = Font(color="7F7F7F", italic=True)
        for col in (2, 3, 7):
            ws_t.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 4, 5):
            ws_t.cell(row=r, column=col).alignment = Alignment(vertical="top")
        ws_t.row_dimensions[r].height = 45
    last = len(task_list) + 1
    ws_t.column_dimensions["F"].hidden = True
    ws_t.cell(row=last + 2, column=3, value="Score").font = Font(bold=True)
    ws_t.cell(row=last + 2, column=4, value=f'=COUNTIF(E2:E{last},"✅*")&" / {len(task_list)}"').font = Font(bold=True, size=13)
    ws_t.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=['LEFT(E2,1)="✅"'], font=Font(color="006100"),
                                                               fill=PatternFill("solid", fgColor="C6EFCE")))
    ws_t.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=['LEFT(E2,1)="❌"'], font=Font(color="9C0006"),
                                                               fill=PatternFill("solid", fgColor="FFC7CE")))
    ws_t.freeze_panes = "A2"

    # ---------------- Orders_Solved (helper columns used by the answer key)
    ws_s = wb.create_sheet("Orders_Solved")
    cols = list(df.columns) + ["Category", "Line Revenue", "Month", "Order Size", "unit_cost", "Weekday #"]
    header(ws_s, 1, cols, [10, 12, 11, 12, 8, 15, 12, 17, 11, 10, 9, 10, 9, 16, 13, 10, 11, 10, 10])
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            c = ws_s.cell(row=i, column=j, value=v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
            if j == 2:
                c.number_format = "yyyy-mm-dd"
        ws_s.cell(row=i, column=14, value=f"=VLOOKUP(J{i},Products!$A$2:$C$49,3,FALSE)")
        ws_s.cell(row=i, column=15, value=f"=K{i}*L{i}-M{i}")
        ws_s.cell(row=i, column=16, value=f'=TEXT(B{i},"yyyy-mm")')
        ws_s.cell(row=i, column=17, value=f'=IF(O{i}>=3000,"Large","Small")')
        ws_s.cell(row=i, column=18, value=f"=INDEX(Products!$F$2:$F$49,MATCH(J{i},Products!$A$2:$A$49,0))")
        ws_s.cell(row=i, column=19, value=f"=WEEKDAY(B{i},2)")  # Monday=1 … Sunday=7
    ws_s.freeze_panes = "A2"

    # ---------------- Answer Key
    ws_a = wb.create_sheet("Answer Key")
    ws_a["A1"] = "Answer Key — one working approach per task (formulas point at Orders_Solved)"
    ws_a["A1"].font = Font(bold=True, size=14, color=NAVY)
    header(ws_a, 3, ["#", "Formula result", "Formula used (see cell B)"], [4, 18, 90])
    for r, (tid, _, _, kind, _, formula) in enumerate(task_list, start=4):
        ws_a.cell(row=r, column=1, value=tid)
        c = ws_a.cell(row=r, column=2, value=formula)
        c.number_format = "0.0%" if kind == "pct" else ("#,##0.00" if kind == "num" else "General")
        text_cell(ws_a, r, 3, formula.replace(R, "").replace("_xlfn.", "")).font = Font(name="Consolas", size=9, color="404040")
    # helper tables for tasks 14 and 17
    ws_a["H3"], ws_a["I3"] = "city", "valid revenue"
    ws_a["K3"], ws_a["L3"] = "weekday", "order lines"
    for c in ("H3", "I3", "K3", "L3"):
        ws_a[c].font, ws_a[c].fill = HFONT, HFILL
    I_, O_, CITY = rng("I"), rng("O"), rng("D")
    for k, city in enumerate(sorted(df.city.unique()), start=4):
        ws_a.cell(row=k, column=8, value=city)
        ws_a.cell(row=k, column=9, value=f'=SUMIFS({O_},{CITY},H{k},{I_},"Delivered")+SUMIFS({O_},{CITY},H{k},{I_},"Shipped")')
    for k, day in enumerate(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"], start=4):
        ws_a.cell(row=k, column=11, value=day)
        # WEEKDAY(...,2): Monday=1 … Sunday=7 (locale-independent, unlike TEXT(date,"dddd"))
        ws_a.cell(row=k, column=12, value=f'=COUNTIF({rng("S")},{k - 3})')
    ws_a.column_dimensions["H"].width = 14
    ws_a.column_dimensions["I"].width = 16
    ws_a.column_dimensions["K"].width = 12
    ws_a.column_dimensions["L"].width = 12

    wb.save(OUT)
    return df, exp


if __name__ == "__main__":
    df, exp = build()
    print(f"Wrote {OUT.name}: {len(df):,} order lines, {len(exp)} tasks")
    for k, v in exp.items():
        print(f"  Task {k:>2}: {v:,.4f}" if not isinstance(v, str) else f"  Task {k:>2}: {v}")
