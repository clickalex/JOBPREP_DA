"""Excel workbook: structure, hidden answers, and (slow) real formula evaluation of the Answer Key."""
import re

import openpyxl
import pytest

from conftest import ROOT

XLSX = ROOT / "practice" / "excel" / "ShopKart_Excel_Practice.xlsx"
EXPECTED = {1: 1349, 2: 8039818.85, 3: 6820604.80, 4: 738239.90, 5: 104, 6: 2581, 7: 601.90, 8: 10392,
            9: 1655441.20, 10: 819, 11: "Face Wash", 12: 0.1547, 13: 0.3749, 14: "Delhi", 15: 0.5197,
            16: 97266.55, 17: "Friday", 18: 3077.89, 19: 631, 20: 12355.65}


def close(a, b):
    if isinstance(b, str):
        return str(a).strip() == b
    return abs(float(a) - b) <= max(0.006, abs(b) * 1e-4) if b > 1 else abs(float(a) - b) <= 0.0006


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(XLSX)


def test_sheets(wb):
    for name in ["Tasks", "Orders", "Products", "Answer Key", "Orders_Solved"]:
        assert name in wb.sheetnames


def test_hidden_expected_answers(wb):
    ws = wb["Tasks"]
    got = {ws.cell(r, 1).value: ws.cell(r, 6).value for r in range(2, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)}
    assert set(got) == set(EXPECTED)
    for k, v in EXPECTED.items():
        assert close(got[k], v), (k, got[k], v)
    assert ws.column_dimensions["F"].hidden


def test_modern_functions_have_xlfn_prefix(wb):
    """Excel needs post-2007 functions stored as _xlfn.NAME or it shows #NAME? until the cell is re-entered."""
    modern = ("MAXIFS", "MINIFS", "XLOOKUP", "IFS", "TEXTJOIN", "CONCAT", "UNIQUE", "FILTER", "SORT", "LET")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f" and isinstance(c.value, str):
                    bare = re.findall(r"(?<![A-Z_.])(" + "|".join(modern) + r")\(", c.value.upper())
                    assert not bare, f"{ws.title}!{c.coordinate}: {c.value}"


@pytest.mark.slow
def test_answer_key_formulas_evaluate_to_expected(wb):
    """Evaluate the whole workbook with the `formulas` engine (≈2 min) and compare Answer Key results."""
    formulas = pytest.importorskip("formulas")
    sol = formulas.ExcelModel().loads(str(XLSX)).finish().calculate()
    vals = {k.upper(): v for k, v in sol.items()}
    ka = wb["Answer Key"]
    checked = 0
    for r in range(1, ka.max_row + 1):
        task = ka.cell(r, 1).value
        if str(task).isdigit() and int(task) in EXPECTED:
            key = next(k for k in vals if k.endswith(f"ANSWER KEY'!B{r}"))
            v = vals[key].value[0, 0]
            assert close(v, EXPECTED[int(task)]), (task, v)
            checked += 1
    assert checked == 20
